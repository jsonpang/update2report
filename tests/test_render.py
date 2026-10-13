from pathlib import Path

from update2report.render_html import render_html
from update2report.render_excel import render_excel

def test_html_is_single_file_and_excel_has_three_sheets(tmp_path):
    payload = {'date':'2026-10-13','tone':'internal','summary':'One item done.','updates':[],'stats':{'projects':{},'overdue':[],'no_updates':[]}}
    html = render_html(payload)
    assert '<html' in html.lower() and 'One item done.' in html
    target = tmp_path / 'report.xlsx'
    render_excel(payload, target)
    from openpyxl import load_workbook
    book = load_workbook(target, read_only=True)
    assert book.sheetnames == ['Summary', 'Details', 'Risks']
import json

from update2report import service

def test_cached_sample_generates_offline_without_llm(monkeypatch, tmp_path):
    sample_dir = tmp_path / 'samples'
    (sample_dir / 'cache').mkdir(parents=True)
    (sample_dir / 'cache' / 'offline.json').write_text(json.dumps({
        'updates': [{'person':'Ari','project':'Atlas','done':['Shipped'],'next':[],'blockers':[],'due_dates':[],'confidence':0.9}],
        'summaries': {'internal':'Atlas work moved forward.'}
    }), encoding='utf-8')
    csv_file = tmp_path / 'tasks.csv'
    csv_file.write_text('task,owner,project,due,status\nShip,Ari,Atlas,2026-10-12,done\n', encoding='utf-8')
    monkeypatch.setattr(service, 'SAMPLE_DIR', sample_dir)
    monkeypatch.setattr(service, 'OUT_DIR', tmp_path / 'out')
    monkeypatch.delenv('LLM_BASE_URL', raising=False)
    monkeypatch.delenv('LLM_MODEL', raising=False)
    summary, html_path, excel_path, stats, tokens = service.generate_report('cached notes', str(csv_file), 'internal', 'offline', report_date=__import__('datetime').date(2026,10,13))
    assert summary == 'Atlas work moved forward.'
    assert Path(html_path).exists() and Path(excel_path).exists()
    assert stats['projects']['Atlas']['completed'] == 1
    assert tokens == 0
