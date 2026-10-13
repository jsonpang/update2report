from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def render_excel(payload, path):
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.append(['Northwind Studio Weekly Update', payload['date']])
    summary.append(['Project', 'Completed', 'Blockers'])
    for name, values in payload['stats']['projects'].items(): summary.append([name, values['completed'], values['blockers']])
    summary.append(['Summary', payload['summary']])
    details = book.create_sheet('Details')
    details.append(['Person','Project','Done','Next','Blockers','Due dates','Confidence'])
    for item in payload['updates']:
        details.append([item.person,item.project,'; '.join(item.done),'; '.join(item.next),'; '.join(item.blockers),'; '.join(item.due_dates),item.confidence])
    risks = book.create_sheet('Risks')
    risks.append(['Type','Task or person','Owner','Project','Due','Status'])
    for task in payload['stats']['overdue']:
        risks.append(['Overdue',task.get('task',''),task.get('owner',''),task.get('project',''),task.get('due',''),task.get('status','')])
    for person in payload['stats']['no_updates']: risks.append(['No update',person,'','','',''])
    if payload['stats'].get('unassigned_blockers'): risks.append(['Unassigned blockers','','','','',payload['stats']['unassigned_blockers']])
    for sheet in book.worksheets:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='263D78')
        for column in sheet.columns:
            width = min(55, max(12, max(len(str(cell.value or '')) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    book.save(path)
